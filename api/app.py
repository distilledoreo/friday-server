"""Desktop Assistant API: one authenticated front door for the local backends.

    /chat    -> llama.cpp  (OpenAI chat-completions body; streaming supported)
    /search  -> SearXNG    (GET: trimmed JSON results; POST: the Android app's web_search contract)
    /fetch   -> Crawl4AI   (page as markdown)
    /transcribe, /speak, /voice -> Parakeet v3 speech-to-text and Kokoro text-to-speech, on this machine
    /extract -> text from attached documents (PDF, Word, text/code); page images for scanned PDFs

Every route except /health needs `Authorization: Bearer $ASSISTANT_API_TOKEN`.
/v1/chat/completions and /v1/models are aliases so OpenAI-style clients work as-is.
"""

import asyncio
import base64
import io
import re
import ipaddress
import os
import secrets
import socket
import threading
import time
from urllib.parse import unquote, urlparse

import httpx
from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse, Response, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, Field

import prompt_cache

LLM_URL = os.environ.get("LLM_URL", "http://127.0.0.1:8080")
SEARXNG_URL = os.environ.get("SEARXNG_URL", "http://127.0.0.1:8888")
CRAWL4AI_URL = os.environ.get("CRAWL4AI_URL", "http://127.0.0.1:11235")
CRAWL4AI_TOKEN = os.environ["CRAWL4AI_API_TOKEN"]
API_TOKEN = os.environ["ASSISTANT_API_TOKEN"]
VOICE_MODELS = os.environ.get("VOICE_MODELS", os.path.expanduser("~/voice-models"))

app = FastAPI(title="Desktop Assistant API", docs_url=None, redoc_url=None, openapi_url=None)
bearer = HTTPBearer(auto_error=False)
# No read timeout for the LLM: long generations are normal.
llm = httpx.AsyncClient(base_url=LLM_URL, timeout=httpx.Timeout(10.0, read=None))
http = httpx.AsyncClient(timeout=httpx.Timeout(10.0, read=120.0))


def require_token(creds: HTTPAuthorizationCredentials | None = Depends(bearer)) -> None:
    if creds is None or not secrets.compare_digest(creds.credentials, API_TOKEN):
        raise HTTPException(401, "invalid or missing token", headers={"WWW-Authenticate": "Bearer"})


auth = [Depends(require_token)]


@app.get("/health")
async def health():
    async def up(client: httpx.AsyncClient, url: str) -> bool:
        try:
            return (await client.get(url, timeout=3.0)).status_code < 500
        except httpx.HTTPError:
            return False

    status = {
        "llm": await up(llm, "/health"),
        "search": await up(http, f"{SEARXNG_URL}/healthz"),
        "fetch": await up(http, f"{CRAWL4AI_URL}/health"),
    }
    return JSONResponse({"ok": all(status.values()), **status}, status_code=200 if all(status.values()) else 503)


# ---- chat -----------------------------------------------------------------

@app.post("/chat", dependencies=auth)
@app.post("/v1/chat/completions", dependencies=auth, include_in_schema=False)
async def chat(request: Request):
    from imagegen.privacy import prepare_chat
    if request.headers.get("X-Assistant-Prime") == "1":
        # The app's new-chat opening, sent ahead of the first message; nothing is generated.
        if request.headers.get("X-Assistant-Incognito") == "1":
            return JSONResponse({"primed": "ignored"})
        return JSONResponse({"primed": openings.prime(prompt_cache.stabilize(await request.json()))})
    body = await prepare_chat(await request.json(), request.headers.get("X-Assistant-Incognito") == "1", prompt_cache, openings)
    req = llm.build_request("POST", "/v1/chat/completions", json=body)
    try:
        upstream = await llm.send(req, stream=True)
    except httpx.HTTPError as e:
        raise HTTPException(502, f"LLM unavailable: {e.__class__.__name__}")

    if not body.get("stream"):
        content = await upstream.aread()
        await upstream.aclose()
        return Response(content, status_code=upstream.status_code, media_type="application/json")

    async def relay():
        try:
            async for chunk in upstream.aiter_raw():
                yield chunk
        finally:
            await upstream.aclose()

    return StreamingResponse(relay(), status_code=upstream.status_code, media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/v1/models", dependencies=auth, include_in_schema=False)
async def models():
    r = await llm.get("/v1/models")
    return JSONResponse(r.json(), status_code=r.status_code)


# ---- search ---------------------------------------------------------------

async def searxng(params: dict) -> dict:
    try:
        r = await http.get(f"{SEARXNG_URL}/search", params={**params, "format": "json"})
        r.raise_for_status()
    except httpx.HTTPError as e:
        raise HTTPException(502, f"search unavailable: {e.__class__.__name__}")
    return r.json()


@app.get("/search", dependencies=auth)
async def search(q: str, n: int = 8, categories: str | None = None, time_range: str | None = None):
    params = {"q": q}
    if categories:
        params["categories"] = categories
    if time_range:
        params["time_range"] = time_range
    data = await searxng(params)
    results = [
        {"title": x.get("title"), "url": x.get("url"), "snippet": x.get("content"), "engine": x.get("engine")}
        for x in data.get("results", [])[: max(1, min(n, 30))]
    ]
    return {"query": q, "results": results, "answers": data.get("answers", [])}


class WebSearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=500)
    limit: int = Field(5, ge=1, le=10)
    fetch_pages: bool = False


PAGE_FETCH_LIMIT = 2
PAGE_TEXT_CHARS = 4000


@app.post("/search", dependencies=auth)
async def web_search(req: WebSearchRequest):
    """The Android app's web_search contract (same shape as desktop/search_service in that repo)."""
    data = await searxng({"q": req.query.strip()})
    hits = [x for x in data.get("results", []) if x.get("url")][: req.limit]

    async def page(url: str) -> tuple[str | None, str | None]:
        try:
            check_public_url(url)
            doc = await crawl(url)
        except HTTPException as e:
            return None, str(e.detail)
        if not doc["success"]:
            return None, doc["error"] or "page fetch failed"
        return doc["markdown"][:PAGE_TEXT_CHARS], None

    pages = await asyncio.gather(*(page(h["url"]) for h in hits[:PAGE_FETCH_LIMIT])) if req.fetch_pages else []
    results = []
    for i, h in enumerate(hits):
        text, error = pages[i] if i < len(pages) else (None, None)
        item = {"title": h.get("title") or "", "url": h["url"], "snippet": h.get("content") or "",
                "score": h.get("score"), "page_text": text}
        if error:
            item["page_error"] = error
        results.append(item)
    return {"provider": "searxng", "fetched": any(p[0] for p in pages), "results": results}


# ---- fetch ----------------------------------------------------------------

# Drop page chrome and prune low-value blocks so fit_markdown is mostly the article.
CRAWL_CONFIG = {"type": "CrawlerRunConfig", "params": {
    "cache_mode": "bypass",
    "excluded_tags": ["nav", "header", "footer", "aside", "form"],
    "markdown_generator": {"type": "DefaultMarkdownGenerator", "params": {
        "content_filter": {"type": "PruningContentFilter", "params": {"threshold": 0.48}},
        "options": {"ignore_images": True},
    }},
}}


class FetchRequest(BaseModel):
    url: str
    max_chars: int = Field(20000, ge=100, le=200000)


def check_public_url(url: str) -> None:
    """Refuse non-http(s) URLs and anything resolving to a private/local address,
    so fetching can't be used to reach services on this machine or the LAN."""
    parts = urlparse(url)
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise HTTPException(400, "url must be http(s) with a host")
    try:
        infos = socket.getaddrinfo(parts.hostname, parts.port or 443)
    except socket.gaierror:
        raise HTTPException(400, "host does not resolve")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise HTTPException(400, "url points at a private or local address")


async def crawl(url: str) -> dict:
    """One page through Crawl4AI, as {url, success, status_code, title, markdown, error}."""
    try:
        r = await http.post(f"{CRAWL4AI_URL}/crawl", headers={"Authorization": f"Bearer {CRAWL4AI_TOKEN}"},
                            json={"urls": [url], "crawler_config": CRAWL_CONFIG})
        r.raise_for_status()
    except httpx.HTTPError as e:
        raise HTTPException(502, f"fetch unavailable: {e.__class__.__name__}")
    result = r.json()["results"][0]
    md = result.get("markdown") or ""
    if isinstance(md, dict):
        md = md.get("fit_markdown") or md.get("raw_markdown") or ""
    return {
        "url": result.get("url", url),
        "success": result.get("success", False),
        "status_code": result.get("status_code"),
        "title": (result.get("metadata") or {}).get("title"),
        "markdown": md,
        "error": result.get("error_message"),
    }


@app.post("/fetch", dependencies=auth)
async def fetch(req: FetchRequest):
    check_public_url(req.url)
    doc = await crawl(req.url)
    md = doc.pop("markdown")
    return {**doc, "markdown": md[: req.max_chars], "truncated": len(md) > req.max_chars}


# ---- voice ----------------------------------------------------------------

class VoiceModels:
    """Parakeet TDT 0.6B v3 (int8) and Kokoro v1.0, loaded once in the background and run on the CPU."""

    def __init__(self):
        self.asr = None
        self.tts = None
        self.error: str | None = None
        self.ready = threading.Event()
        self.asr_lock = threading.Lock()
        self.tts_lock = threading.Lock()

    def load(self):
        try:
            import onnx_asr
            from kokoro_onnx import Kokoro

            self.asr = onnx_asr.load_model(
                "nemo-parakeet-tdt-0.6b-v3", f"{VOICE_MODELS}/parakeet-tdt-0.6b-v3-int8", quantization="int8")
            self.tts = Kokoro(f"{VOICE_MODELS}/kokoro-v1.0/kokoro-v1.0.onnx", f"{VOICE_MODELS}/kokoro-v1.0/voices-v1.0.bin")
            self.error = None
        except Exception as e:  # models missing, Storage drive not mounted, ...
            self.error = f"{e.__class__.__name__}: {e}"
        finally:
            self.ready.set()

    def require(self):
        if not self.ready.wait(timeout=60) or self.error or self.asr is None:
            raise HTTPException(503, f"voice models unavailable: {self.error or 'still loading'}")


voice = VoiceModels()
threading.Thread(target=voice.load, name="voice-load", daemon=True).start()

MAX_AUDIO_BYTES = 16_000 * 2 * 120  # two minutes of 16 kHz PCM16


def decode_audio(body: bytes, content_type: str):
    """WAV, or raw little-endian PCM16 mono with `audio/L16; rate=N`. Returns float32 at 16 kHz."""
    import numpy as np
    import soundfile as sf

    if body[:4] == b"RIFF":
        data, rate = sf.read(io.BytesIO(body), dtype="float32", always_2d=True)
        samples = data.mean(axis=1)
    else:
        rate = 16_000
        for part in content_type.split(";"):
            key, _, value = part.strip().partition("=")
            if key.lower() == "rate" and value.isdigit():
                rate = int(value)
        samples = np.frombuffer(body[: len(body) // 2 * 2], dtype="<i2").astype(np.float32) / 32768.0
    if rate != 16_000 and len(samples):
        positions = np.arange(0, len(samples), rate / 16_000)
        samples = np.interp(positions, np.arange(len(samples)), samples).astype(np.float32)
    return samples


@app.get("/voice", dependencies=auth)
async def voice_status():
    voice.require()
    return {"stt": "parakeet-tdt-0.6b-v3", "tts": "kokoro-v1.0", "voices": sorted(voice.tts.get_voices())}


@app.post("/transcribe", dependencies=auth)
async def transcribe(request: Request):
    body = await request.body()
    if not body:
        raise HTTPException(400, "send audio as WAV or audio/L16")
    if len(body) > MAX_AUDIO_BYTES:
        raise HTTPException(413, "audio longer than two minutes")
    voice.require()
    started = time.monotonic()
    try:
        samples = decode_audio(body, request.headers.get("content-type", ""))
    except Exception:
        raise HTTPException(400, "could not decode audio")
    if len(samples) < 1_600:  # under 0.1 s
        return {"text": "", "duration_ms": len(samples) // 16, "elapsed_ms": 0}

    def run():
        with voice.asr_lock:
            return voice.asr.recognize(samples, sample_rate=16_000)

    text = (await run_in_threadpool(run)).strip()
    return {"text": text, "duration_ms": len(samples) // 16, "elapsed_ms": int((time.monotonic() - started) * 1000)}


class SpeakRequest(BaseModel):
    text: str = Field(min_length=1, max_length=1_000)
    voice: str = "af_heart"
    speed: float = Field(1.0, ge=0.5, le=2.0)
    lang: str = "en-us"


@app.post("/speak", dependencies=auth)
async def speak(req: SpeakRequest):
    voice.require()
    if req.voice not in voice.tts.get_voices():
        raise HTTPException(400, f"unknown voice '{req.voice}'")

    def run():
        import numpy as np
        import soundfile as sf

        with voice.tts_lock:
            try:
                samples, rate = voice.tts.create(req.text, voice=req.voice, speed=req.speed, lang=req.lang)
            except ValueError as e:  # e.g. text with nothing pronounceable, like "-"
                raise HTTPException(422, str(e))
        out = io.BytesIO()
        sf.write(out, np.clip(samples, -1.0, 1.0), rate, format="WAV", subtype="PCM_16")
        return out.getvalue()

    return Response(await run_in_threadpool(run), media_type="audio/wav")


# ---- documents --------------------------------------------------------------

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_DOCUMENT_CHARS = 60_000
MAX_SCANNED_PAGES = 4
TEXT_SUFFIXES = {
    "txt", "md", "markdown", "csv", "tsv", "json", "jsonl", "xml", "yaml", "yml", "toml", "ini", "log", "html", "htm",
    "py", "kt", "kts", "java", "js", "ts", "tsx", "jsx", "c", "h", "cpp", "hpp", "cs", "go", "rs", "rb", "php", "swift",
    "sh", "bash", "zsh", "sql", "css", "scss", "gradle", "properties", "srt", "vtt", "tex", "rtf",
}


def _clip(text: str) -> tuple[str, bool]:
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return (text[:MAX_DOCUMENT_CHARS], True) if len(text) > MAX_DOCUMENT_CHARS else (text, False)


def extract_document(data: bytes, name: str, mime: str) -> dict:
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if data[:5] == b"%PDF-" or suffix == "pdf" or mime == "application/pdf":
        import pypdf
        import pypdfium2

        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages) if t)
        result = {"kind": "pdf", "pages": len(pages), "images": []}
        if len(re.sub(r"\s+", "", text)) < 40 * max(1, min(len(pages), 3)):
            # No usable text layer (a scan): send page images for the vision model instead.
            pdf = pypdfium2.PdfDocument(data)
            for i in range(min(len(pdf), MAX_SCANNED_PAGES)):
                page = pdf[i]
                scale = 1280 / max(page.get_width(), page.get_height())
                image = page.render(scale=max(scale, 0.5)).to_pil().convert("RGB")
                out = io.BytesIO()
                image.save(out, format="JPEG", quality=85)
                result["images"].append(base64.b64encode(out.getvalue()).decode())
            result["note"] = (
                f"No text layer; sent {len(result['images'])} page image(s)"
                + (f" of {len(pdf)}" if len(pdf) > MAX_SCANNED_PAGES else "")
                + "."
            )
            text = ""
        result["text"], result["truncated"] = _clip(text)
        return result
    if data[:2] == b"PK" and (suffix == "docx" or "wordprocessingml" in mime):
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text, truncated = _clip("\n".join(parts))
        return {"kind": "docx", "text": text, "truncated": truncated}
    if suffix in TEXT_SUFFIXES or mime.startswith("text/") or mime in ("application/json", "application/xml"):
        if b"\x00" in data[:4096]:
            raise HTTPException(415, "that file looks binary, not text")
        raw = data.decode("utf-8", errors="replace")
        if suffix in ("html", "htm") or mime == "text/html":
            raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
            raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        text, truncated = _clip(raw)
        return {"kind": "text", "text": text, "truncated": truncated}
    raise HTTPException(415, f"can't read {suffix or mime or 'this'} files yet; try PDF, Word, or a text file")


@app.post("/extract", dependencies=auth)
async def extract(request: Request):
    data = await request.body()
    if not data:
        raise HTTPException(400, "send the file as the request body")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise HTTPException(413, "file larger than 25 MB")
    name = unquote(request.headers.get("x-filename", "document"))[:200]
    mime = request.headers.get("content-type", "").split(";")[0].strip().lower()
    try:
        result = await run_in_threadpool(extract_document, data, name, mime)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"couldn't read {name}: {e.__class__.__name__}")
    return {"name": name, "chars": len(result["text"]), **result}


# ---- persistent workspace extension ----
from pathlib import Path as WorkspacePath
from workspace.integration import enable as enable_workspace
# ---- local image GPU handoff ----
from imagegen.integration import prepare as prepare_images, enable as enable_images
raw_llm = llm
llm, gpu_gate = prepare_images(raw_llm)
openings = prompt_cache.OpeningSnapshots(raw_llm, gpu_gate)
app.on_event("startup")(openings.start)
app.on_event("shutdown")(openings.stop)
workspace_store = enable_workspace(app, auth, llm, search, fetch, FetchRequest,
    WorkspacePath(__file__).resolve().parent.parent / "workspace-data")

image_manager = enable_images(app, auth, workspace_store, raw_llm, gpu_gate,
    WorkspacePath(__file__).resolve().parent / "imagegen" / "config.json")

from memory.integration import enable as enable_memory
memory_store = enable_memory(app, auth,
    WorkspacePath(__file__).resolve().parent.parent / 'workspace-data' / 'memory',
    llm, gpu_gate)


# ---- FRIDAY isolated cloud agent ----
from agent.integration import enable as enable_friday_agent
from pathlib import Path as FridayAgentPath
friday_agent_store = enable_friday_agent(app, auth,
    FridayAgentPath(__file__).resolve().parent.parent / "workspace-data" / "agent", search, gpu_gate)
